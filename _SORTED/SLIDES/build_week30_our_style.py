from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

source = r"C:\Users\mrmic\Downloads\PL12 - Week 30.pptx"
out = r"C:\Users\mrmic\Desktop\OB-1\_SORTED\SLIDES\PL12 - Week 30 - OUR STYLE.pptx"

prs = Presentation(source)

# Add subtle brand ribbon on all slides
for slide in prs.slides:
    ribbon = slide.shapes.add_shape(
        1, Inches(11.6), Inches(0.0), Inches(1.73), Inches(0.38)
    )
    ribbon.fill.solid()
    ribbon.fill.fore_color.rgb = RGBColor(16, 44, 84)
    ribbon.line.fill.background()
    tf = ribbon.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = "NET-PET-AI"
    run.font.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER

# Add vocabulary page
layout = prs.slide_layouts[6]  # blank
vocab_slide = prs.slides.add_slide(layout)

bg = vocab_slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.33), Inches(7.5))
bg.fill.solid()
bg.fill.fore_color.rgb = RGBColor(245, 248, 252)
bg.line.fill.background()

header = vocab_slide.shapes.add_shape(
    1, Inches(0), Inches(0), Inches(13.33), Inches(1.15)
)
header.fill.solid()
header.fill.fore_color.rgb = RGBColor(16, 44, 84)
header.line.fill.background()

box = vocab_slide.shapes.add_textbox(Inches(0.7), Inches(0.2), Inches(10), Inches(0.8))
tf = box.text_frame
p = tf.paragraphs[0]
r = p.add_run()
r.text = "Vocabulary Bank"
r.font.size = Pt(38)
r.font.bold = True
r.font.color.rgb = RGBColor(255, 255, 255)

sub = vocab_slide.shapes.add_textbox(
    Inches(0.72), Inches(1.32), Inches(12), Inches(0.5)
)
stf = sub.text_frame
sp = stf.paragraphs[0]
sr = sp.add_run()
sr.text = "Unit 8 - The Media | Key words for speaking and listening"
sr.font.size = Pt(19)
sr.font.bold = True
sr.font.color.rgb = RGBColor(16, 44, 84)

vocab = [
    ("self-esteem", "how confident and valuable you feel about yourself"),
    ("express", "to show your thoughts or feelings clearly"),
    ("cyberbullying", "using the internet to hurt or threaten others"),
    ("anonymous", "without showing your real name or identity"),
    ("gossip", "unofficial talk about other people's private lives"),
    ("embarrassment", "the shy or uncomfortable feeling after a mistake"),
    ("trustworthy", "reliable and likely to give correct information"),
    ("body image", "how you think and feel about your body"),
]

start_y = 1.95
row_h = 0.63
for i, (word, meaning) in enumerate(vocab):
    y = start_y + i * row_h
    card = vocab_slide.shapes.add_shape(
        1, Inches(0.7), Inches(y), Inches(11.9), Inches(0.5)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card.line.color.rgb = RGBColor(210, 220, 234)

    wb = vocab_slide.shapes.add_textbox(
        Inches(0.95), Inches(y + 0.08), Inches(3.2), Inches(0.35)
    )
    wtf = wb.text_frame
    wp = wtf.paragraphs[0]
    wr = wp.add_run()
    wr.text = word
    wr.font.size = Pt(20)
    wr.font.bold = True
    wr.font.color.rgb = RGBColor(16, 44, 84)

    mb = vocab_slide.shapes.add_textbox(
        Inches(3.55), Inches(y + 0.1), Inches(8.7), Inches(0.35)
    )
    mtf = mb.text_frame
    mp = mtf.paragraphs[0]
    mr = mp.add_run()
    mr.text = meaning
    mr.font.size = Pt(16)
    mr.font.color.rgb = RGBColor(57, 74, 99)

footer = vocab_slide.shapes.add_textbox(
    Inches(0.7), Inches(7.0), Inches(12), Inches(0.3)
)
ft = footer.text_frame
fp = ft.paragraphs[0]
fr = fp.add_run()
fr.text = "Tip: Ask students to make one original sentence with each word."
fr.font.size = Pt(12)
fr.font.italic = True
fr.font.color.rgb = RGBColor(90, 104, 125)

# Move the vocabulary slide to position 9 (1-based)
slide_ids = prs.slides._sldIdLst
slides = list(slide_ids)
new_id = slides[-1]
slide_ids.remove(new_id)
insert_idx = 8 if len(slides) >= 8 else len(slides)
slide_ids.insert(insert_idx, new_id)

os.makedirs(os.path.dirname(out), exist_ok=True)
prs.save(out)
print(out)
